#!/usr/bin/env -S uv run --quiet
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0", "lxml>=5.0"]
# ///
"""Real PPTX regressions for the text-size gate and Japanese run fonts."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "slide-creator" / "scripts"
sys.path.insert(0, str(SCRIPTS))


class ReadabilityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="slide-text-test-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def box(self, text="検査用の本文", size=22, shapes=None, name="body"):
        shapes = self.slide.shapes if shapes is None else shapes
        box = shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.name = name
        box.text = text
        if size is not None:
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        return box

    def verify(self, expected=0, exceptions=None):
        deck = self.root / "test.pptx"
        self.prs.save(deck)
        command = [sys.executable, str(SCRIPTS / "verify_deck.py"), str(deck)]
        if exceptions is not None:
            allowlist = self.root / "text-exceptions.json"
            allowlist.write_text(json.dumps(exceptions), encoding="utf-8")
            command += ["--font-size-exceptions", str(allowlist)]
        original = deck.read_bytes()
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        self.assertEqual(deck.read_bytes(), original, "verification must not mutate the PPTX")
        return result.stdout + result.stderr

    def test_body_floor_and_mixed_runs(self):
        self.box(size=16)
        self.verify()
        tiny = self.box(text="小さすぎる本文", size=15.99, name="tiny")
        paragraph = tiny.text_frame.paragraphs[0]
        paragraph.add_run().text = " large run"
        paragraph.runs[1].font.size = Pt(24)
        output = self.verify(expected=1)
        self.assertIn("15.99pt", output)
        self.assertIn("tiny", output)
        self.assertIn("小さすぎる本文", output)

    def test_paragraph_list_and_default_inheritance(self):
        box = self.box(size=None)
        self.verify()  # built-in master otherStyle is 18pt
        paragraph = box.text_frame.paragraphs[0]
        paragraph.font.size = Pt(14)
        self.assertIn("14pt", self.verify(expected=1))
        paragraph.font.size = None
        lst_style = box.text_frame._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lstStyle")
        level = OxmlElement("a:lvl2pPr")
        rpr = OxmlElement("a:defRPr")
        rpr.set("sz", "1500")
        level.append(rpr)
        lst_style.append(level)
        paragraph.level = 1
        self.assertIn("15pt", self.verify(expected=1))

    def test_layout_and_master_placeholder_inheritance(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        slide.shapes.title.text = "Inherited title"
        body = slide.placeholders[1]
        body.text = "Inherited body"
        layout_body = slide.slide_layout.placeholders[1]
        layout_body.text_frame.paragraphs[0].font.size = Pt(13)
        self.assertIn("13pt", self.verify(expected=1))
        layout_body.text_frame.paragraphs[0].font.size = None
        master_body = next(ph for ph in slide.slide_layout.slide_master.placeholders
                           if ph._element.ph_type == body._base_placeholder._base_placeholder._element.ph_type)
        master_body.text_frame.paragraphs[0].font.size = Pt(12)
        self.assertIn("12pt", self.verify(expected=1))
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
        self.verify()

    def test_groups_tables_fields_and_autofit(self):
        group = self.slide.shapes.add_group_shape()
        self.box(size=12, shapes=group.shapes, name="group-body")
        table_shape = self.slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(4), Inches(1))
        table_shape.name = "table-body"
        cell = table_shape.table.cell(0, 0)
        cell.text = "小さな表"
        cell.text_frame.paragraphs[0].font.size = Pt(13)
        field_box = self.box(text="", size=None, name="field-body")
        field = OxmlElement("a:fld")
        field.set("id", "{8DE39BE0-7D00-4FD8-AB6E-DA0961B32980}")
        field.set("type", "slidenum")
        rpr = OxmlElement("a:rPr")
        rpr.set("sz", "1100")
        field.append(rpr)
        text = OxmlElement("a:t")
        text.text = "1"
        field.append(text)
        field_box.text_frame.paragraphs[0]._p.append(field)
        shrink = self.box(size=20, name="shrunk-body")
        body_pr = shrink.text_frame._txBody.bodyPr
        for tag in ("a:spAutoFit", "a:normAutofit", "a:noAutofit"):
            element = body_pr.find(tag, body_pr.nsmap)
            if element is not None:
                body_pr.remove(element)
        auto = OxmlElement("a:normAutofit")
        auto.set("fontScale", "75000")
        body_pr.append(auto)
        output = self.verify(expected=1)
        for name in ("group-body", "table-body", "field-body", "shrunk-body"):
            self.assertIn(name, output)
        self.assertIn("15pt", output)

    def test_specific_exceptions_do_not_waive_other_body_text(self):
        self.box(text="Source: Example", size=12, name="scid:s001.source")
        entry = {"slide": 1, "shape": "scid:s001.source", "role": "source",
                 "reason": "Bibliographic credit only; explanation stays in body"}
        output = self.verify(exceptions={"exceptions": [entry]})
        self.assertIn("EXEMPT", output)
        self.assertIn(entry["reason"], output)
        self.box(size=15, name="small-body")
        self.assertIn("small-body", self.verify(expected=1, exceptions={"exceptions": [entry]}))
        for patch in ({"shape": "missing"}, {"reason": ""}, {"role": "body"}, {"shpae": "typo"}):
            self.verify(expected=1, exceptions={"exceptions": [{**entry, **patch}]})
        self.verify(expected=1, exceptions={"exceptions": [entry, entry]})

    def test_missing_size_fails_instead_of_assuming_a_safe_font(self):
        self.box(size=None)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for root in (self.prs._element, self.slide.slide_layout.slide_master._element):
            for element in root.findall(".//a:defRPr", ns):
                element.attrib.pop("sz", None)
        self.assertIn("unresolved font size", self.verify(expected=1))

    def test_master_style_uses_the_paragraph_level_and_title_role(self):
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        slide.shapes.title.text = "Title inherits titleStyle"
        body = slide.placeholders[1]
        body.text = "Body inherits level two"
        body.text_frame.paragraphs[0].level = 1
        master = slide.slide_layout.slide_master._element
        master.find("p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr", ns).set("sz", "1400")
        master.find("p:txStyles/p:bodyStyle/a:lvl2pPr/a:defRPr", ns).set("sz", "1500")
        output = self.verify(expected=1)
        self.assertIn("14pt", output)
        self.assertIn("15pt", output)

    def test_empty_paragraph_and_end_paragraph_properties_are_not_visible_text(self):
        box = self.box(size=22)
        end = OxmlElement("a:endParaRPr")
        end.set("sz", "1000")
        box.text_frame.paragraphs[0]._p.append(end)
        box.text_frame.add_paragraph().font.size = Pt(8)
        self.verify()

    def test_duplicate_shape_names_cannot_make_an_exception_ambiguous(self):
        for _ in range(2):
            self.box(size=12, name="source")
        output = self.verify(expected=1, exceptions={"exceptions": [
            {"slide": 1, "shape": "source", "role": "source", "reason": "Citation only"}
        ]})
        self.assertIn("matched 2 text shapes", output)

    def test_math_is_checked_at_its_base_size(self):
        from add_equation import build_a14_wrapper
        box = self.box(text="", size=None)
        para = box.text_frame.paragraphs[0]
        para.font.size = Pt(14)
        para._p.append(build_a14_wrapper(
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:r><m:t>x=1</m:t></m:r></m:oMath>', display=True))
        self.assertIn("14pt", self.verify(expected=1))
        para.font.size = Pt(24)
        self.verify()

    def test_saved_autofit_inherits_and_local_no_autofit_overrides_it(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text = "Inherited saved scaling"
        layout_body = slide.slide_layout.placeholders[1]
        auto = OxmlElement("a:normAutofit")
        auto.set("fontScale", "40000")
        layout_body.text_frame._txBody.bodyPr.append(auto)
        self.assertIn("12.8pt", self.verify(expected=1))
        body.text_frame._txBody.bodyPr.append(OxmlElement("a:noAutofit"))
        self.verify()

    def test_japanese_run_font_preserves_hyperlinks_and_other_properties(self):
        from text_style import set_japanese_font
        box = self.box()
        run = box.text_frame.paragraphs[0].runs[0]
        run.font.bold = True
        run.hyperlink.address = "https://example.com/"
        set_japanese_font(run, "Yu Gothic")
        set_japanese_font(run, "Hiragino Kaku Gothic ProN")
        self.prs.save(self.root / "font.pptx")
        reopened = Presentation(self.root / "font.pptx").slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        rpr = reopened._r.rPr
        for script in ("latin", "ea", "cs"):
            font = rpr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{script}")
            self.assertEqual(len(font), 1)
            self.assertEqual(font[0].get("typeface"), "Hiragino Kaku Gothic ProN")
        self.assertTrue(reopened.font.bold)
        self.assertEqual(reopened.font.size.pt, 22)
        self.assertEqual(reopened.hyperlink.address, "https://example.com/")
        self.assertLess(list(rpr).index(font[0]), list(rpr).index(rpr.hlinkClick))


if __name__ == "__main__":
    unittest.main()
