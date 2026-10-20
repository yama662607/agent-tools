#!/usr/bin/env -S uv run --quiet
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=0.6.23", "lxml>=5.0"]
# ///
"""Print-size audit of a one-page .pptx poster. Read-only; never renders.

Reports (all at PRINT size; use --scale N when the page was built at 1/N):
  page      page size in mm and ISO A-size detection
  fonts     per-run font sizes against poster floors (FAIL < 18 pt,
            WARN < 24 pt; references/credits shapes may sit at 18-23 pt);
            autofit fontScale is applied, inherited sizes are flagged
  pictures  effective dpi at the placed size, after cropping
            (FAIL < 150, WARN < 300)
  bounds    shapes outside the page (FAIL) or inside the margin (WARN)
  overflow  text-height heuristic per text box (approximate: CJK = 1 em,
            Latin = 0.5 em; confirm in a PowerPoint PDF)
  overlap   partial overlap between top-level blocks (full containment such
            as text on a card is normal and not reported)

Exit 1 when any FAIL is found. Usage:
  poster_check.py poster.pptx [--scale N] [--margin-mm 15] [--slide 1] [--json]
"""
import argparse
import json
import math
import re
import sys
import unicodedata

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

EMU_PER_MM = 36000
EMU_PER_IN = 914400
PAPER_MM = {"A0": (841, 1189), "A1": (594, 841), "A2": (420, 594),
            "A3": (297, 420), "A4": (210, 297)}
FAIL_PT, WARN_PT = 18.0, 24.0
FAIL_DPI, WARN_DPI = 150.0, 300.0
SMALL_TEXT_NAME_RE = re.compile(
    r"ref|acknowledg|credit|funding|謝辞|参考|文献|affil|所属|contact|source",
    re.IGNORECASE)
LINE_FACTOR = 1.2  # single-spaced line height / font size


def detect_paper(w_mm, h_mm, tol=3.0):
    for name, (a, b) in PAPER_MM.items():
        if abs(w_mm - a) <= tol and abs(h_mm - b) <= tol:
            return f"{name} portrait"
        if abs(w_mm - b) <= tol and abs(h_mm - a) <= tol:
            return f"{name} landscape"
    return None


class Box:
    """A leaf shape with absolute page coordinates in EMU."""

    def __init__(self, shape, x, y, w, h, top_level):
        self.shape, self.x, self.y, self.w, self.h = shape, x, y, w, h
        self.top_level = top_level
        self.name = shape.name or "<unnamed>"


def walk(shapes, tf=None, top_level=True):
    """Yield Box for every leaf shape, mapping group child coordinates."""
    for sh in shapes:
        if sh.left is None or sh.width is None:
            continue
        if tf is None:
            x, y, w, h = sh.left, sh.top, sh.width, sh.height
        else:
            ox, oy, sx, sy = tf
            x, y = ox + sh.left * sx, oy + sh.top * sy
            w, h = sh.width * sx, sh.height * sy
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = sh._element.grpSpPr.find(qn("a:xfrm"))
            ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
            sx = w / int(ch_ext.get("cx")) if int(ch_ext.get("cx")) else 1
            sy = h / int(ch_ext.get("cy")) if int(ch_ext.get("cy")) else 1
            child_tf = (x - int(ch_off.get("x")) * sx,
                        y - int(ch_off.get("y")) * sy, sx, sy)
            yield from walk(sh.shapes, child_tf, False)
            continue
        yield Box(sh, x, y, w, h, top_level)


def run_size_pt(run_el, lst_sizes, para_lvl):
    """(size_pt, inherited) for an a:r / a:fld element."""
    rpr = run_el.find(qn("a:rPr"))
    if rpr is not None and rpr.get("sz"):
        return int(rpr.get("sz")) / 100, False
    if para_lvl in lst_sizes:
        return lst_sizes[para_lvl], True
    return 18.0, True  # PowerPoint default when nothing is stated


def text_info(box):
    """Per-paragraph runs for the box: list of dicts, plus fontScale."""
    sh = box.shape
    if not sh.has_text_frame:
        return [], 1.0
    body = sh._element.find(".//" + qn("p:txBody"))
    if body is None:
        return [], 1.0
    bpr = body.find(qn("a:bodyPr"))
    scale = 1.0
    if bpr is not None:
        na = bpr.find(qn("a:normAutofit"))
        if na is not None and na.get("fontScale"):
            scale = int(na.get("fontScale")) / 100000
    lst_sizes = {}
    lst = body.find(qn("a:lstStyle"))
    if lst is not None:
        for lvl in range(1, 10):
            d = lst.find(f"{qn('a:lvl%dpPr' % lvl)}/{qn('a:defRPr')}")
            if d is not None and d.get("sz"):
                lst_sizes[lvl - 1] = int(d.get("sz")) / 100
    paras = []
    for p in body.findall(qn("a:p")):
        ppr = p.find(qn("a:pPr"))
        lvl = int(ppr.get("lvl", 0)) if ppr is not None else 0
        runs = []
        for r in p.findall(qn("a:r")) + p.findall(qn("a:fld")):
            text = "".join(t.text or "" for t in r.findall(qn("a:t")))
            if not text.strip():
                continue
            size, inh = run_size_pt(r, lst_sizes, lvl)
            runs.append({"text": text, "size": size * scale, "inherited": inh})
        paras.append({"runs": runs, "ppr": ppr})
    return paras, scale


def char_em(ch):
    return 1.0 if unicodedata.east_asian_width(ch) in "WF" else 0.5


def estimate_text_height_pt(box, paras, scale_n):
    """Approximate rendered text height in print points, or None to skip."""
    sh = box.shape
    body = sh._element.find(".//" + qn("p:txBody"))
    bpr = body.find(qn("a:bodyPr"))
    if bpr is not None and bpr.get("wrap") == "none":
        return None
    ins = lambda k, d: int(bpr.get(k, d)) if bpr is not None else d  # noqa: E731
    avail_pt = (box.w - ins("lIns", 91440) - ins("rIns", 91440)) / 12700 * scale_n
    total = 0.0
    for para in paras:
        runs = para["runs"]
        if not runs:
            total += 18.0 * LINE_FACTOR
            continue
        ppr = para["ppr"]
        indent_pt = 0.0
        if ppr is not None and ppr.get("marL"):
            indent_pt = int(ppr.get("marL")) / 12700 * scale_n
        width_pt = sum(char_em(c) * r["size"] for r in runs for c in r["text"])
        size = max(r["size"] for r in runs)
        lines = max(1, math.ceil(width_pt / max(avail_pt - indent_pt, 1)))
        spacing = LINE_FACTOR
        if ppr is not None:
            pct = ppr.find(f"{qn('a:lnSpc')}/{qn('a:spcPct')}")
            if pct is not None:
                spacing = LINE_FACTOR * int(pct.get("val")) / 100000
        total += lines * size * spacing
        if ppr is not None:
            for tag in ("a:spcBef", "a:spcAft"):
                pts = ppr.find(f"{qn(tag)}/{qn('a:spcPts')}")
                if pts is not None:
                    total += int(pts.get("val")) / 100
    return total


def picture_dpi(box, scale_n):
    sh = box.shape
    if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
        return None
    try:
        px_w, px_h = sh.image.size
    except Exception:
        return None
    vis_w = px_w * (1 - sh.crop_left - sh.crop_right)
    vis_h = px_h * (1 - sh.crop_top - sh.crop_bottom)
    w_in = box.w / EMU_PER_IN * scale_n
    h_in = box.h / EMU_PER_IN * scale_n
    if w_in <= 0 or h_in <= 0:
        return None
    return min(vis_w / w_in, vis_h / h_in), (px_w, px_h)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("pptx")
    ap.add_argument("--scale", type=float, default=1.0, metavar="N",
                    help="page is built at 1/N of the print size (default 1)")
    ap.add_argument("--margin-mm", type=float, default=15.0,
                    help="print margin to keep clear, in mm (default 15)")
    ap.add_argument("--slide", type=int, default=1,
                    help="1-based slide holding the poster (default 1)")
    ap.add_argument("--json", action="store_true", help="machine-readable output")
    args = ap.parse_args()
    if args.scale <= 0:
        sys.exit("error: --scale must be greater than zero")

    prs = Presentation(args.pptx)
    if not 1 <= args.slide <= len(prs.slides):
        sys.exit(f"error: --slide {args.slide} out of range ({len(prs.slides)} slides)")
    slide = prs.slides[args.slide - 1]
    n = args.scale
    page_w, page_h = prs.slide_width, prs.slide_height
    w_mm, h_mm = page_w / EMU_PER_MM * n, page_h / EMU_PER_MM * n
    paper = detect_paper(w_mm, h_mm)

    findings = []  # (level, area, where, message)

    def add(level, area, where, msg):
        findings.append({"level": level, "area": area, "where": where, "message": msg})

    if len(prs.slides) != 1:
        add("WARN", "page", "-", f"{len(prs.slides)} slides; a poster is normally one slide "
            f"(checking slide {args.slide})")
    if paper is None:
        add("WARN", "page", "-", f"print size {w_mm:.0f} x {h_mm:.0f} mm is not an ISO A size")

    boxes = list(walk(slide.shapes))
    margin = args.margin_mm * EMU_PER_MM / n  # in page EMU
    page_area = page_w * page_h
    max_font = 0.0
    char_by_size = {}

    for b in boxes:
        full_bleed = b.w * b.h >= 0.9 * page_area
        # bounds
        if not full_bleed:
            over = max(-b.x, -b.y, b.x + b.w - page_w, b.y + b.h - page_h)
            if over > 0.5 * EMU_PER_MM / n:
                add("FAIL", "bounds", b.name, f"extends {over / EMU_PER_MM * n:.1f} mm outside the page")
            elif (b.x < margin or b.y < margin or b.x + b.w > page_w - margin
                  or b.y + b.h > page_h - margin):
                add("WARN", "bounds", b.name, f"inside the {args.margin_mm:g} mm margin")
        # fonts + overflow
        paras, _ = text_info(b)
        small_ok = bool(SMALL_TEXT_NAME_RE.search(b.name))
        flagged_fail = flagged_warn = False
        for para in paras:
            for r in para["runs"]:
                pt = r["size"] * n
                max_font = max(max_font, pt)
                char_by_size[round(pt)] = char_by_size.get(round(pt), 0) + len(r["text"].strip())
                tag = " (inherited, assumed)" if r["inherited"] else ""
                if pt + 1e-6 < FAIL_PT and not flagged_fail:
                    add("FAIL", "fonts", b.name, f"{pt:g} pt < {FAIL_PT:g} pt floor{tag}: {r['text'].strip()[:40]!r}")
                    flagged_fail = True
                elif pt + 1e-6 < WARN_PT and not small_ok and not flagged_warn and pt >= FAIL_PT:
                    add("WARN", "fonts", b.name, f"{pt:g} pt < {WARN_PT:g} pt (ok only for references/credits){tag}: {r['text'].strip()[:40]!r}")
                    flagged_warn = True
        if paras and any(p["runs"] for p in paras):
            est = estimate_text_height_pt(b, paras, n)
            if est is not None:
                body = b.shape._element.find(".//" + qn("p:txBody")).find(qn("a:bodyPr"))
                tb = (int(body.get("tIns", 45720)) + int(body.get("bIns", 45720))) / 12700 * n
                avail = b.h / 12700 * n - tb
                grows = body.find(qn("a:spAutoFit")) is not None
                if avail > 0 and est > avail * 1.05 and grows:
                    add("WARN", "overflow", b.name, f"resize-shape-to-fit box will grow: text ~{est:.0f} pt vs stored {avail:.0f} pt (may collide with neighbours)")
                elif avail > 0 and est > avail * 1.05:
                    add("FAIL", "overflow", b.name, f"estimated text {est:.0f} pt tall vs box {avail:.0f} pt ({est / avail:.0%})")
                elif avail > 0 and est > avail * 0.9 and not grows:
                    add("WARN", "overflow", b.name, f"estimated text fills {est / avail:.0%} of the box")
        # pictures
        pic = picture_dpi(b, n)
        if pic is not None:
            dpi, (pw, ph) = pic
            if dpi < FAIL_DPI:
                add("FAIL", "pictures", b.name, f"{dpi:.0f} dpi at print size ({pw}x{ph} px) < {FAIL_DPI:g}")
            elif dpi < WARN_DPI:
                add("WARN", "pictures", b.name, f"{dpi:.0f} dpi at print size ({pw}x{ph} px) < {WARN_DPI:g}")

    if max_font and max_font < 72:
        add("WARN", "fonts", "-", f"largest text is {max_font:g} pt; a poster title is normally >= 72 pt")

    # overlap between top-level blocks (partial only)
    def thin(b):
        return min(b.w, b.h) < 0.02 * EMU_PER_IN * 3 / n  # lines / rules
    top = [b for b in boxes if b.top_level and not thin(b) and b.w * b.h < 0.9 * page_area]
    tol = 2 * EMU_PER_MM / n
    for i, a in enumerate(top):
        for c in top[i + 1:]:
            ix = min(a.x + a.w, c.x + c.w) - max(a.x, c.x)
            iy = min(a.y + a.h, c.y + c.h) - max(a.y, c.y)
            if ix <= tol or iy <= tol:
                continue
            def contains(p, q):
                return (p.x - tol <= q.x and p.y - tol <= q.y
                        and p.x + p.w + tol >= q.x + q.w and p.y + p.h + tol >= q.y + q.h)
            if contains(a, c) or contains(c, a):
                continue
            frac = ix * iy / min(a.w * a.h, c.w * c.h)
            if frac > 0.02:
                add("WARN", "overlap", f"{a.name} / {c.name}",
                    f"partial overlap {ix / EMU_PER_MM * n:.0f} x {iy / EMU_PER_MM * n:.0f} mm ({frac:.0%} of the smaller block)")

    result = {
        "file": args.pptx, "slide": args.slide, "scale": n,
        "page_mm": [round(w_mm, 1), round(h_mm, 1)], "paper": paper,
        "shapes": len(boxes), "largest_font_pt": max_font or None,
        "chars_by_font_pt": dict(sorted(char_by_size.items())),
        "fail": sum(f["level"] == "FAIL" for f in findings),
        "warn": sum(f["level"] == "WARN" for f in findings),
        "findings": findings,
    }
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        note = f" (built at 1/{n:g})" if n != 1 else ""
        print(f"{args.pptx}  slide {args.slide}: {w_mm:.0f} x {h_mm:.0f} mm = "
              f"{paper or 'non-ISO'}{note}; {len(boxes)} shapes")
        if char_by_size:
            print("chars by font size (pt): " + ", ".join(f"{k}:{v}" for k, v in sorted(char_by_size.items())))
        if findings:
            print(f"{'LEVEL':5}  {'AREA':8}  {'SHAPE':28}  MESSAGE")
            for f in findings:
                print(f"{f['level']:5}  {f['area']:8}  {f['where'][:28]:28}  {f['message']}")
        print(f"RESULT: {result['fail']} FAIL, {result['warn']} WARN")
    sys.exit(1 if result["fail"] else 0)


if __name__ == "__main__":
    main()
